"""High-performance batch data importer with dynamic schema detection."""

from __future__ import annotations
import csv
import json
import os
import sys
import time
from pathlib import Path
from typing import List, Dict, Any, Optional, Iterator, Callable

from src.schema import UserRecord, match_column, clean_col_name
from src.normalization import normalize_phone, is_valid_phone
from src.database import SQLiteRepository


class DataImporter:
    """Handles parsing, normalizing, and streaming batch ingestion into SQLite."""

    def __init__(
        self,
        repository: SQLiteRepository,
        batch_size: int = 25000,
        progress_callback: Optional[Callable[[int, int, int, int, float], None]] = None,
    ):
        self.repo = repository
        self.batch_size = batch_size
        self.progress_callback = progress_callback

        # Stats
        self.total_processed = 0
        self.valid_phones = 0
        self.invalid_phones = 0
        self.duplicate_phone_count = 0

    def import_csv(self, file_path: str | Path, source_name: Optional[str] = None) -> Dict[str, Any]:
        """Import records from a CSV file in streaming batches."""
        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(f"Source file not found: {path}")

        source_label = source_name or path.name
        start_time = time.time()

        self.repo.connect()
        self.repo.init_schema()
        self.repo.prepare_for_bulk_import()

        seen_phones = set()
        batch: List[UserRecord] = []

        with open(path, "r", encoding="utf-8", errors="ignore") as f:
            reader = csv.reader(f)
            header = next(reader, None)
            if not header:
                return {"status": "error", "message": "Empty file"}

            col_map = self._build_column_mapping(header)
            phone_idx = col_map.get("phone")

            if phone_idx is None:
                raise ValueError(
                    f"No phone column could be detected in header: {header[:10]}..."
                )

            for row in reader:
                self.total_processed += 1
                if not row or len(row) <= phone_idx:
                    self.invalid_phones += 1
                    continue

                raw_phone = row[phone_idx].strip()
                norm_phone = normalize_phone(raw_phone)

                if not norm_phone or not is_valid_phone(norm_phone):
                    self.invalid_phones += 1
                    continue

                self.valid_phones += 1
                if norm_phone in seen_phones:
                    self.duplicate_phone_count += 1
                else:
                    seen_phones.add(norm_phone)

                # Extract mapped fields
                rec = self._row_to_record(row, header, col_map, norm_phone, raw_phone, source_label)
                batch.append(rec)

                if len(batch) >= self.batch_size:
                    with self.repo.conn:  # type: ignore
                        self.repo.insert_batch(batch)
                    batch = []

                    if self.progress_callback:
                        elapsed = time.time() - start_time
                        self.progress_callback(
                            self.total_processed,
                            self.valid_phones,
                            self.invalid_phones,
                            self.duplicate_phone_count,
                            elapsed,
                        )

            # Insert remaining records
            if batch:
                with self.repo.conn:  # type: ignore
                    self.repo.insert_batch(batch)

        print("\nFinalizing database indexes and optimizations...")
        self.repo.finalize_bulk_import()

        elapsed = time.time() - start_time
        return {
            "status": "success",
            "source": source_label,
            "total_processed": self.total_processed,
            "valid_phones": self.valid_phones,
            "invalid_phones": self.invalid_phones,
            "duplicate_phones": self.duplicate_phone_count,
            "elapsed_seconds": elapsed,
            "rate_per_sec": self.total_processed / elapsed if elapsed > 0 else 0,
        }

    def import_excel(self, file_path: str | Path, source_name: Optional[str] = None) -> Dict[str, Any]:
        """Import records from an Excel (.xlsx) file in streaming mode."""
        import openpyxl

        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(f"Excel file not found: {path}")

        source_label = source_name or path.name
        start_time = time.time()

        self.repo.connect()
        self.repo.init_schema()
        self.repo.prepare_for_bulk_import()

        seen_phones = set()
        batch: List[UserRecord] = []

        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows_iter = ws.iter_rows(values_only=True)

            try:
                header_row = next(rows_iter)
            except StopIteration:
                continue

            if not header_row or not any(header_row):
                continue

            header = [str(c).strip() if c is not None else "" for c in header_row]
            col_map = self._build_column_mapping(header)
            phone_idx = col_map.get("phone")

            if phone_idx is None:
                continue

            for row_tuple in rows_iter:
                if not row_tuple or not any(row_tuple):
                    continue

                self.total_processed += 1
                row = [str(c).strip() if c is not None else "" for c in row_tuple]

                if len(row) <= phone_idx:
                    self.invalid_phones += 1
                    continue

                raw_phone = row[phone_idx]
                norm_phone = normalize_phone(raw_phone)

                if not norm_phone or not is_valid_phone(norm_phone):
                    self.invalid_phones += 1
                    continue

                self.valid_phones += 1
                if norm_phone in seen_phones:
                    self.duplicate_phone_count += 1
                else:
                    seen_phones.add(norm_phone)

                rec = self._row_to_record(row, header, col_map, norm_phone, raw_phone, f"{source_label}:{sheet_name}")
                batch.append(rec)

                if len(batch) >= self.batch_size:
                    with self.repo.conn:  # type: ignore
                        self.repo.insert_batch(batch)
                    batch = []

                    if self.progress_callback:
                        elapsed = time.time() - start_time
                        self.progress_callback(
                            self.total_processed,
                            self.valid_phones,
                            self.invalid_phones,
                            self.duplicate_phone_count,
                            elapsed,
                        )

            if batch:
                with self.repo.conn:  # type: ignore
                    self.repo.insert_batch(batch)
                batch = []

        wb.close()
        self.repo.finalize_bulk_import()

        elapsed = time.time() - start_time
        return {
            "status": "success",
            "source": source_label,
            "total_processed": self.total_processed,
            "valid_phones": self.valid_phones,
            "invalid_phones": self.invalid_phones,
            "duplicate_phones": self.duplicate_phone_count,
            "elapsed_seconds": elapsed,
            "rate_per_sec": self.total_processed / elapsed if elapsed > 0 else 0,
        }

    def _build_column_mapping(self, header: List[str]) -> Dict[str, int]:
        """Match header column names to standard fields."""
        col_map: Dict[str, int] = {}
        for idx, col in enumerate(header):
            matched = match_column(col)
            if matched and matched not in col_map:
                col_map[matched] = idx
        return col_map

    def _row_to_record(
        self,
        row: List[str],
        header: List[str],
        col_map: Dict[str, int],
        norm_phone: str,
        raw_phone: str,
        source_label: str,
    ) -> UserRecord:
        """Construct UserRecord from row and mapping."""
        def get_val(field_name: str) -> Optional[str]:
            idx = col_map.get(field_name)
            if idx is not None and idx < len(row):
                v = row[idx].strip()
                return v if v and v.lower() not in ("nan", "none", "null") else None
            return None

        # Build extra data JSON for unmapped or e-commerce columns
        extra = {}
        mapped_indices = set(col_map.values())
        for idx, col_name in enumerate(header):
            if idx not in mapped_indices and idx < len(row):
                val = row[idx].strip()
                if val and val.lower() not in ("nan", "none", "null"):
                    extra[col_name] = val

        extra_json = json.dumps(extra, ensure_ascii=False) if extra else None

        return UserRecord(
            phone_normalized=norm_phone,
            phone_raw=raw_phone,
            name=get_val("name"),
            email=get_val("email"),
            city=get_val("city"),
            state=get_val("state"),
            address=get_val("address"),
            gender=get_val("gender"),
            dob=get_val("dob"),
            bank=get_val("bank"),
            account_type=get_val("account_type"),
            income=get_val("income"),
            pan=get_val("pan"),
            profession_1=get_val("profession_1"),
            profession_2=get_val("profession_2"),
            mobile_alerts=get_val("mobile_alerts"),
            extra_data=extra_json,
            source_file=source_label,
        )
